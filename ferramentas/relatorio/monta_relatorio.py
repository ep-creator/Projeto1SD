#!/usr/bin/env python3
"""Monta o relatório do Projeto1SD em .docx (e .pdf).

Fontes:
  docs/relatorio/partes/*.md          capa, introdução, visão geral, ... (em ordem)
  testes/<bloco>/tabela_verdade.md    funcionamento, tabela verdade, equações e mapas-K
  testes/<bloco>/<bloco>_circuito.png print do esquemático (opcional)
  testes/<bloco>/<bloco>_simulacao.png print da waveform (opcional)
  docs/relatorio/figuras/*            demais figuras

Diretivas nas partes:
  <!-- BLOCO: nome -->            insere o bloco (título nível 3)
  <!-- BLOCO: nome | 0 -->        insere o bloco sem título próprio
  <!-- FIGURA: caminho | legenda -->
  <!-- PAGINA -->  <!-- SUMARIO -->

Figura que não existe vira um quadro amarelo "[inserir figura: caminho]".

Uso (na raiz do repositório):
  python3 ferramentas/relatorio/monta_relatorio.py
Saída: docs/relatorio/Relatorio_Projeto1SD.docx e .pdf
Precisa de pandoc, python-docx e (para sumário e PDF) LibreOffice.
"""
import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
PARTES = os.path.join(ROOT, "docs", "relatorio", "partes")
SAIDA = os.path.join(ROOT, "docs", "relatorio", "Relatorio_Projeto1SD.docx")

LARG_MAX_CM = 16.0
ALT_MAX_CM = 21.0

cont = {"tab": 0, "fig": 0}
pendentes = []


# ------------------------------------------------------------------ markdown

def div(style, text):
    return f'::: {{custom-style="{style}"}}\n{text}\n:::\n'


def raw(xml):
    return f"```{{=openxml}}\n{xml}\n```\n"


PAGINA = raw('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
SUMARIO = raw(
    '<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Sumário</w:t></w:r></w:p>'
    '<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/>'
    '</w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>(sumário: atualize o campo)</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p></w:sdtContent></w:sdt>')


def tamanho_img(path):
    try:
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
    except Exception:
        return f"width={LARG_MAX_CM}cm"
    if h / w * LARG_MAX_CM > ALT_MAX_CM:
        return f"height={ALT_MAX_CM}cm"
    return f"width={LARG_MAX_CM}cm"


def figura(rel, legenda):
    cont["fig"] += 1
    n = cont["fig"]
    path = None
    base, _ = os.path.splitext(os.path.join(ROOT, rel))
    for ext in (".png", ".jpg", ".jpeg", ".PNG", ".JPG"):
        if os.path.exists(base + ext):
            path = base + ext
            break
    if path:
        corpo = div("Figura", f"![]({path.replace(os.sep, '/')}){{{tamanho_img(path)}}}")
    else:
        pendentes.append(rel)
        corpo = div("Pendente", f"[inserir figura: {rel}]")
    return corpo + "\n" + div("Legenda Figura", f"Figura {n} – {legenda}") + "\n"


LINK = re.compile(r"\[(`[^`]+`)\]\([^)]*\)")


def secoes(md):
    """divide um tabela_verdade.md em cabeçalho + {título: texto}"""
    partes = re.split(r"^## +(.+)$", md, flags=re.M)
    cab = partes[0]
    d = {}
    for i in range(1, len(partes), 2):
        d[partes[i].strip()] = partes[i + 1].strip()
    return cab, d


def bloco(nome, nivel_titulo=3):
    path = os.path.join(ROOT, "testes", nome, "tabela_verdade.md")
    with open(path, encoding="utf-8") as fh:
        cab, sec = secoes(fh.read())
    h = "#" * (nivel_titulo + 1 if nivel_titulo else 4)
    out = []
    if nivel_titulo:
        out.append(f"{'#' * nivel_titulo} `{nome}`\n")
    else:
        out.append(f"<!-- bloco-atual: {nome} -->\n")
    # interface a partir do cabeçalho
    interface = []
    for linha in cab.splitlines():
        m = re.match(r"\* \*\*(Entradas?|Saídas?):\*\* (.*)", linha)
        if m:
            interface.append(f"* **{m.group(1)}:** " + LINK.sub(r"\1", m.group(2)))
        m = re.match(r"\* \*\*Circuito:\*\* \[[^\]]*\]\([^)]*\) \((.*)\)\s*$", linha)
        if m:
            interface.append("* **Usa:** " + LINK.sub(r"\1", m.group(1)))
    out.append(f"{h} Funcionamento\n\n{sec.get('Funcionamento', '')}\n\n" + "\n".join(interface) + "\n")
    for t in ("Tabela verdade", "Equações", "Mapas de Karnaugh"):
        if sec.get(t):
            out.append(f"{h} {t}\n\n{sec[t]}\n")
    out.append(f"{h} Circuito\n\n" + figura(f"testes/{nome}/{nome}_circuito.png", f"Circuito do `{nome}` (lib/{nome}.bdf)"))
    out.append(f"{h} Simulação\n\n" + figura(f"testes/{nome}/{nome}_simulacao.png", f"Simulação do `{nome}` ({nome}.vwf)"))
    if sec.get("Observações"):
        out.append(f"{h} Observações\n\n{sec['Observações']}\n")
    return "\n".join(out)


def expande(md):
    def rep(m):
        cmd, arg = m.group(1).strip(), (m.group(2) or "").strip()
        if cmd == "PAGINA":
            return PAGINA
        if cmd == "SUMARIO":
            return SUMARIO
        if cmd == "BLOCO":
            p = [x.strip() for x in arg.split("|")]
            return bloco(p[0], int(p[1]) if len(p) > 1 else 3)
        if cmd == "FIGURA":
            rel, leg = [x.strip() for x in arg.split("|", 1)]
            return figura(rel, leg)
        return m.group(0)
    return re.sub(r"<!--\s*(PAGINA|SUMARIO|BLOCO|FIGURA)\s*:?\s*(.*?)-->", rep, md)


def numera_e_legenda(md):
    """numera títulos 1-3 e põe 'Tabela N – ...' antes de cada tabela"""
    nums = [0, 0, 0, 0]
    out = []
    linhas = md.split("\n")
    ctx = {1: "", 2: "", 3: "", 4: ""}
    bloco_atual = ""
    rotulo, sub = "", ""
    em_codigo = False
    anterior = ""
    for ln in linhas:
        if ln.startswith("```"):
            em_codigo = not em_codigo
        m = re.match(r"^(#{1,4}) (.*)$", ln) if not em_codigo else None
        if m:
            nv, txt = len(m.group(1)), m.group(2).strip()
            if nv <= 3:
                nums[nv - 1] += 1
                for k in range(nv, 4):
                    nums[k] = 0
                ln = f"{m.group(1)} {'.'.join(str(x) for x in nums[:nv])} {txt}"
                ctx[nv] = txt
                for k in range(nv + 1, 5):
                    ctx[k] = ""
                if nv == 3 and txt.startswith("`"):
                    bloco_atual = txt.strip("`")
                elif nv < 3 or (nv == 3 and not txt.startswith("`")):
                    bloco_atual = ""
            else:
                ctx[4] = txt
            rotulo, sub = "", ""
            if nv == 1:
                out.append(PAGINA)
            out.append(ln)
            anterior = ln
            continue
        mb = re.match(r"^<!-- bloco-atual: (\S+) -->$", ln.strip())
        if mb:
            bloco_atual = mb.group(1)
            continue
        mr = re.match(r"^\*\*([^*]+)\*\*$", ln.strip())
        ms = re.match(r"^\*([^*]+)\*$", ln.strip())
        if not em_codigo and mr:
            rotulo, sub = mr.group(1), ""
            continue
        if not em_codigo and ms:
            sub = ms.group(1)
            continue
        if ln.startswith("|") and not anterior.startswith("|") and not em_codigo:
            cont["tab"] += 1
            secao = ctx[4] or ctx[3] or ctx[2] or ctx[1]
            secao = secao.strip("`")
            if secao == "Mapas de Karnaugh":
                desc = f"mapa de Karnaugh de {rotulo}" if rotulo else "mapa de Karnaugh"
            else:
                desc = secao[0].lower() + secao[1:] if secao else ""
                if rotulo:
                    desc += f" ({rotulo[0].lower() + rotulo[1:]})" if secao else rotulo
            if sub:
                desc += f", {sub}"
            pre = f"`{bloco_atual}`: " if bloco_atual else ""
            out.append("")
            out.append(div("Legenda Tabela", f"Tabela {cont['tab']} – {pre}{desc}"))
        out.append(ln)
        anterior = ln
    return "\n".join(out)


def monta_md():
    partes = sorted(glob.glob(os.path.join(PARTES, "*.md")))
    md = "\n\n".join(open(p, encoding="utf-8").read() for p in partes)
    return numera_e_legenda(expande(md))


# ------------------------------------------------------------------ docx

def set_cell_shading(cell, cor):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:shd")):
        tcPr.remove(e)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def bordas(tbl):
    tblPr = tbl._tbl.tblPr
    for e in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(e)
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "808080")
        b.append(e)
    tblPr.append(b)
    # largura automática
    for e in tblPr.findall(qn("w:tblW")):
        tblPr.remove(e)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "0")
    w.set(qn("w:type"), "auto")
    tblPr.append(w)
    for e in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(e)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "autofit")
    tblPr.append(lay)
    # margens das células
    for e in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(e)
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("left", 70), ("right", 70)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)
    # remove larguras fixas de colunas/células
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc in grid.findall(qn("w:gridCol")):
            gc.attrib.pop(qn("w:w"), None)
    for tc in tbl._tbl.iter(qn("w:tc")):
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is not None:
            for e in tcPr.findall(qn("w:tcW")):
                tcPr.remove(e)


CM = 567  # dxa por cm


def larguras(tbl, tam, mono_cols=()):
    """largura de cada coluna pelo conteúdo; tabela com layout fixo"""
    linhas = [[("".join(t.text or "" for t in tc.iter(qn("w:t")))).strip() for tc in tr.findall(qn("w:tc"))]
              for tr in tbl._tbl.findall(qn("w:tr"))]
    ncol = max(len(r) for r in linhas)
    chars = [0] * ncol
    for i, row in enumerate(linhas):
        for j, txt in enumerate(row):
            if i == 0 and "\\" not in txt:
                n = max((len(w) for w in txt.split()), default=0)
            else:
                n = len(txt)
            chars[j] = max(chars[j], n)
    cw = tam * 0.0353 * 0.56
    ws = [max(0.9, c * cw + 0.45) for c in chars]
    tot = sum(ws)
    if tot > LARG_MAX_CM:
        ws = [w * LARG_MAX_CM / tot for w in ws]
    tblPr = tbl._tbl.tblPr
    for e in tblPr.findall(qn("w:tblW")) + tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(e)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(sum(ws) * CM)))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl._tbl.insert(1, grid)
    for gc in grid.findall(qn("w:gridCol")):
        grid.remove(gc)
    for wc in ws:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(wc * CM)))
        grid.append(gc)
    for tr in tbl._tbl.findall(qn("w:tr")):
        for j, tc in enumerate(tr.findall(qn("w:tc"))[:ncol]):
            tcPr = tc.get_or_add_tcPr()
            e = OxmlElement("w:tcW")
            e.set(qn("w:w"), str(int(ws[j] * CM)))
            e.set(qn("w:type"), "dxa")
            tcPr.insert(0, e)


def repete_cabecalho(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)
    e = OxmlElement("w:cantSplit")
    trPr.append(e)


def estilo(doc, nome):
    alvo = nome.lower().replace(" ", "")
    for st in doc.styles:
        if st.name and st.name.lower().replace(" ", "") == alvo:
            return st
        if getattr(st, "style_id", None) and st.style_id.lower() == alvo:
            return st
    return None


def fonte(st, nome=None, tam=None, negrito=None, cor=None, italico=None):
    if st is None:
        return
    f = st.font
    if nome:
        f.name = nome
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), nome)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            rf.attrib.pop(qn(a), None)
    if tam:
        f.size = Pt(tam)
    if negrito is not None:
        f.bold = negrito
    if italico is not None:
        f.italic = italico
    if cor is not None:
        f.color.rgb = RGBColor.from_string(cor)
        c = st.element.get_or_add_rPr().find(qn("w:color"))
        if c is not None:
            for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
                c.attrib.pop(qn(a), None)


def paragrafo(st, alin=None, antes=None, depois=None, linha=None, keep=None, quebra=None):
    if st is None:
        return
    pf = st.paragraph_format
    if alin is not None:
        pf.alignment = alin
    if antes is not None:
        pf.space_before = Pt(antes)
    if depois is not None:
        pf.space_after = Pt(depois)
    if linha is not None:
        pf.line_spacing = linha
    if keep is not None:
        pf.keep_with_next = keep
    if quebra is not None:
        pf.page_break_before = quebra


def campo_pagina(par):
    r = par.add_run()
    for tipo, txt in (("begin", None), (None, " PAGE "), ("end", None)):
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
            r._r.append(e)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
            r._r.append(e)


def formata(path):
    doc = Document(path)
    TEXTO = "Times New Roman"
    TITULO = "Arial"
    MONO = "Courier New"

    for nome in ("Normal", "Body Text", "First Paragraph", "Compact"):
        fonte(estilo(doc, nome), TEXTO, 12, cor="000000")
    for nome in ("Body Text", "First Paragraph"):
        paragrafo(estilo(doc, nome), WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, 1.5)
    fonte(estilo(doc, "Compact"), TEXTO, 11)
    paragrafo(estilo(doc, "Compact"), None, 0, 0, 1.0)
    for i, tam in ((1, 15), (2, 13), (3, 12), (4, 12)):
        st = estilo(doc, f"Heading {i}")
        fonte(st, TITULO, tam if i < 4 else 11, True, "000000", italico=False)
        paragrafo(st, WD_ALIGN_PARAGRAPH.LEFT, 18 if i < 4 else 10, 6, 1.15, True, False)
    fonte(estilo(doc, "Verbatim Char"), MONO, 10.5)
    st = estilo(doc, "Source Code")
    fonte(st, MONO, 9)
    paragrafo(st, WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 1.0)
    for nome, alin, antes, depois, tam, neg in (
            ("Legenda Tabela", WD_ALIGN_PARAGRAPH.CENTER, 8, 3, 10, False),
            ("Legenda Figura", WD_ALIGN_PARAGRAPH.CENTER, 3, 10, 10, False),
            ("Figura", WD_ALIGN_PARAGRAPH.CENTER, 6, 0, 12, None),
            ("Pendente", WD_ALIGN_PARAGRAPH.CENTER, 6, 0, 12, True),
            ("Capa Instituicao", WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 13, True),
            ("Capa Titulo", WD_ALIGN_PARAGRAPH.CENTER, 0, 6, 18, True),
            ("Capa Subtitulo", WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 12, False),
            ("Capa Equipe", WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 12, False),
            ("Capa Local", WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 12, False)):
        st = estilo(doc, nome)
        fonte(st, TITULO if nome.startswith("Capa") else TEXTO, tam, neg, "000000")
        paragrafo(st, alin, antes, depois, 1.15)
    paragrafo(estilo(doc, "Legenda Tabela"), keep=True)
    paragrafo(estilo(doc, "Figura"), keep=True)
    paragrafo(estilo(doc, "Capa Titulo"), antes=110)
    paragrafo(estilo(doc, "Capa Subtitulo"), depois=70)
    paragrafo(estilo(doc, "Capa Equipe"), depois=16)
    paragrafo(estilo(doc, "Capa Local"), antes=40)
    st = estilo(doc, "TOC Heading")
    fonte(st, TITULO, 15, True, "000000")
    paragrafo(st, WD_ALIGN_PARAGRAPH.CENTER, 0, 12)
    for i in (1, 2, 3):
        st = estilo(doc, f"TOC {i}") or estilo(doc, f"toc {i}")
        fonte(st, TEXTO, 12 if i == 1 else 11, i == 1)

    # quadro amarelo nas figuras pendentes
    for p in doc.paragraphs:
        if p.style.name == "Pendente":
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "FFF2CC")
            pPr.append(shd)
            bdr = OxmlElement("w:pBdr")
            for lado in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{lado}")
                e.set(qn("w:val"), "dashed")
                e.set(qn("w:sz"), "8")
                e.set(qn("w:space"), "8")
                e.set(qn("w:color"), "BF9000")
                bdr.append(e)
            pPr.insert(0, bdr)
            for r in p.runs:
                r.font.size = Pt(11)

    # tabelas
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordas(t)
        ncol = max(len(tr.findall(qn("w:tc"))) for tr in t._tbl.findall(qn("w:tr")))
        tam = 10 if ncol <= 8 else (9 if ncol <= 12 else 8)
        kmap = "\\" in t.rows[0].cells[0].text
        larguras(t, tam)
        for i, row in enumerate(t.rows):
            if i == 0:
                repete_cabecalho(row)
            for j, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i == 0:
                    set_cell_shading(cell, "D9D9D9")
                elif kmap and j == 0:
                    set_cell_shading(cell, "D9D9D9")
                elif kmap and cell.text.strip() == "1":
                    set_cell_shading(cell, "C6EFCE")
                elif kmap and cell.text.strip() == "X":
                    set_cell_shading(cell, "EDEDED")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_with_next = False
                    for r in p.runs:
                        r.font.size = Pt(tam)
                        if i == 0:
                            r.font.bold = True
                    if kmap:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Word atualiza o sumário ao abrir
    st_el = doc.settings.element
    if st_el.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        st_el.append(uf)

    # página A4, margens ABNT, número da página
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin, s.left_margin = Cm(3), Cm(3)
        s.bottom_margin, s.right_margin = Cm(2), Cm(2)
        s.different_first_page_header_footer = True
        fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        campo_pagina(fp)
    doc.save(path)


# ------------------------------------------------------------------ LibreOffice

UNO = r'''
import sys, time, subprocess, uno
from com.sun.star.beans import PropertyValue
def prop(n, v):
    p = PropertyValue(); p.Name = n; p.Value = v; return p
src, pdf = sys.argv[1], sys.argv[2]
proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore",
    "--accept=socket,host=localhost,port=2083;urp;"])
ctx = None
for _ in range(60):
    try:
        local = uno.getComponentContext()
        res = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        ctx = res.resolve("uno:socket,host=localhost,port=2083;urp;StarOffice.ComponentContext")
        break
    except Exception:
        time.sleep(0.5)
desk = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
url = uno.systemPathToFileUrl(src)
doc = desk.loadComponentFromURL(url, "_blank", 0, (prop("Hidden", True),))
idx = doc.getDocumentIndexes()
for i in range(idx.getCount()):
    idx.getByIndex(i).update()
doc.refresh()
idx = doc.getDocumentIndexes()
for i in range(idx.getCount()):
    idx.getByIndex(i).update()
doc.storeToURL(uno.systemPathToFileUrl(pdf), (prop("FilterName", "writer_pdf_Export"),))
doc.close(True)
try:
    desk.terminate()
except Exception:
    pass
proc.wait(timeout=30)
'''


def libreoffice(docx, pdf):
    with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False) as fh:
        fh.write(UNO)
        script = fh.name
    py = "python3"
    for cand in ("/usr/lib/libreoffice/program/python", "/opt/libreoffice/program/python"):
        if os.path.exists(cand):
            py = cand
    env = dict(os.environ)
    env["PYTHONPATH"] = "/usr/lib/libreoffice/program:" + env.get("PYTHONPATH", "")
    try:
        subprocess.run([py, script, docx, pdf], check=True, timeout=240, env=env)
        return True
    except Exception as e:
        print("aviso: LibreOffice não atualizou o sumário/PDF:", e)
        return False
    finally:
        os.unlink(script)
        for lock in glob.glob(os.path.join(os.path.dirname(docx), ".~lock.*#")):
            try:
                os.remove(lock)
            except OSError:
                pass


def main():
    md = monta_md()
    tmp = os.path.join(tempfile.gettempdir(), "relatorio_montado.md")
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write(md)
    subprocess.run(["pandoc", tmp, "-f", "markdown+pipe_tables-implicit_figures", "-t", "docx", "-o", SAIDA],
                   check=True, cwd=ROOT)
    formata(SAIDA)
    pdf = SAIDA[:-5] + ".pdf"
    if "--sem-pdf" not in sys.argv and shutil.which("soffice"):
        libreoffice(SAIDA, pdf)
    print(f"{cont['tab']} tabelas, {cont['fig']} figuras")
    if pendentes:
        print(f"{len(pendentes)} figuras pendentes:")
        for p in pendentes:
            print("  ", p)
    print("gerado:", SAIDA)


if __name__ == "__main__":
    main()
